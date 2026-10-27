from __future__ import annotations

"""
Дублирование регистраций команд в локальный Excel-файл (.xlsx).

Никаких внешних сервисов не требует — файл просто лежит рядом с ботом.
При регистрации добавляется строка, при отмене регистрации строка
удаляется — в таблице всегда остаются только действующие команды.
Открывается в Excel, Numbers, Google Таблицах (импортом) и т.д.

Управляется через .env:
  EXCEL_EXPORT_ENABLED — "1" (по умолчанию) или "0", чтобы отключить
  EXCEL_FILE_PATH      — путь к файлу, по умолчанию "teams.xlsx"
"""
import asyncio
import logging
import os

logger = logging.getLogger(__name__)

_enabled = os.getenv("EXCEL_EXPORT_ENABLED", "1") == "1"
EXCEL_PATH = os.getenv("EXCEL_FILE_PATH", "teams.xlsx")
SHEET_TITLE = "Регистрации"

HEADER = [
    "Telegram ID",
    "Telegram username",
    "Команда",
    "Капитан",
    "MMR",
    "Основной состав",
    "Замены",
    "Дата регистрации",
]

_lock = asyncio.Lock()


def is_enabled() -> bool:
    return _enabled


def _open_or_create_workbook():
    from openpyxl import Workbook, load_workbook

    if os.path.exists(EXCEL_PATH):
        wb = load_workbook(EXCEL_PATH)
        if SHEET_TITLE not in wb.sheetnames:
            ws = wb.create_sheet(SHEET_TITLE)
            ws.append(HEADER)
        return wb, wb[SHEET_TITLE]

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    ws.append(HEADER)

    # немного расширяем колонки, чтобы файл было удобно открыть
    widths = [14, 20, 22, 20, 10, 40, 30, 19]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

    return wb, ws


def _append_registration_sync(
    tg_user_id: int,
    tg_username: str | None,
    team_name: str,
    captain_contact: str,
    mmr: str,
    players_info: str,
    substitutes_info: str,
    registered_at: str,
) -> bool:
    try:
        wb, ws = _open_or_create_workbook()
        ws.append(
            [
                tg_user_id,
                f"@{tg_username}" if tg_username else "",
                team_name,
                captain_contact,
                mmr,
                players_info,
                substitutes_info,
                registered_at,
            ]
        )
        wb.save(EXCEL_PATH)
        return True
    except Exception:
        logger.exception("Не удалось записать регистрацию в Excel-файл")
        return False


def _remove_registration_sync(tg_user_id: int) -> bool:
    try:
        wb, ws = _open_or_create_workbook()
        id_col = HEADER.index("Telegram ID") + 1

        found_row = None
        for row in ws.iter_rows(min_row=2):
            cell = row[id_col - 1]
            if cell.value == tg_user_id:
                found_row = cell.row
                break

        if found_row is not None:
            ws.delete_rows(found_row)
            wb.save(EXCEL_PATH)
            return True
        return False
    except Exception:
        logger.exception("Не удалось удалить регистрацию из Excel-файла")
        return False


async def append_registration(
    tg_user_id: int,
    tg_username: str | None,
    team_name: str,
    captain_contact: str,
    mmr: str,
    players_info: str,
    substitutes_info: str,
    registered_at: str,
) -> bool:
    if not _enabled:
        return False
    async with _lock:
        return await asyncio.to_thread(
            _append_registration_sync,
            tg_user_id,
            tg_username,
            team_name,
            captain_contact,
            mmr,
            players_info,
            substitutes_info,
            registered_at,
        )


async def remove_registration(tg_user_id: int) -> bool:
    """Полностью удаляет строку команды из файла (используется при отмене регистрации)."""
    if not _enabled:
        return False
    async with _lock:
        return await asyncio.to_thread(_remove_registration_sync, tg_user_id)